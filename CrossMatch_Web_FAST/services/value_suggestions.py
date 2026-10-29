"""Private, on-demand, disk-indexed suggestions for exact-value condition text.

Each source column is indexed only on first use. Subsequent keystrokes use an
indexed SQLite prefix query (not a scan of the uploaded workbook/CSV). The
session-scoped temporary cache is removed with the CrossMatch session.
"""
from __future__ import annotations

import hashlib
import json
import os
import sqlite3
import threading
import uuid
from pathlib import Path
from typing import Iterator

import pandas as pd
from engine import crossmatch_engine as engine
from services.session_manager import session_dir

_INDEX_LOCK = threading.RLock()
MAX_VALUE_LENGTH = 500  # Equal to the Rule Builder text input maxlength.


def _column_values(plan: engine.InputPlan, column: str) -> Iterator[str]:
    """Read one CSV column, or stream the selected Excel worksheet."""
    if plan.path.suffix.casefold() == '.csv':
        reader = pd.read_csv(
            plan.path, sep=plan.csv_sep, encoding=plan.csv_encoding,
            skiprows=plan.header_row, header=0, names=plan.headers,
            usecols=[column], dtype=str, keep_default_na=False,
            na_filter=False, low_memory=False, chunksize=20000, engine='c',
        )
        for batch in reader:
            yield from batch[column]
        return

    from openpyxl import load_workbook
    wb = load_workbook(plan.path, read_only=True, data_only=True, keep_links=False)
    try:
        ws = wb[plan.sheet or wb.sheetnames[0]]
        column_pos = plan.headers.index(column)
        for row in ws.iter_rows(min_row=plan.header_row + 2, values_only=True):
            if column_pos < len(row):
                yield engine._excel_cell_text(row[column_pos])
    finally:
        wb.close()


def _cache_file(session_id: str, side: str, file_info: dict, column: str) -> Path:
    identity = json.dumps([
        'autocomplete-v2', side, file_info.get('sha256'), file_info.get('server_path'),
        file_info.get('sheet'), file_info.get('header_row'), column,
    ], ensure_ascii=False)
    digest = hashlib.sha256(identity.encode('utf-8')).hexdigest()
    return session_dir(session_id) / 'working' / f'value_suggestions_{side}_{digest}.sqlite'


def _build_index(db_file: Path, plan: engine.InputPlan, column: str) -> None:
    db_file.parent.mkdir(parents=True, exist_ok=True)
    temp = db_file.with_name(db_file.name + '.' + uuid.uuid4().hex + '.tmp')
    try:
        with sqlite3.connect(temp) as db:
            # Safe for disposable caches: incomplete builds never become the final DB.
            db.execute('PRAGMA synchronous=OFF')
            db.execute('PRAGMA journal_mode=OFF')
            db.execute('CREATE TABLE suggestions (folded TEXT PRIMARY KEY, value TEXT NOT NULL) WITHOUT ROWID')
            values = []
            for raw in _column_values(plan, column):
                text = str(raw if raw is not None else '').strip()
                if not text or len(text) > MAX_VALUE_LENGTH:
                    continue
                values.append((text.casefold(), text))
                if len(values) >= 5000:
                    db.executemany('INSERT OR IGNORE INTO suggestions VALUES (?,?)', values)
                    values.clear()
            if values:
                db.executemany('INSERT OR IGNORE INTO suggestions VALUES (?,?)', values)
            db.execute('CREATE TABLE metadata (count INTEGER NOT NULL)')
            db.execute('INSERT INTO metadata VALUES ((SELECT COUNT(*) FROM suggestions))')
            db.commit()
        os.replace(temp, db_file)
    finally:
        temp.unlink(missing_ok=True)


def lookup_suggestions_for_file(session_id: str, side: str, file_info: dict,
                                column: str, query: str, limit: int = 10) -> dict:
    """Return actual cell values from one selected column, not header names.

    Fast ordered prefix lookup first; if fewer than ten results match, try
    substring search so typing part of a middle word also works. An empty
    query shows examples from the selected uploaded column on focus.
    """
    search = query.strip().casefold()
    if len(search) > MAX_VALUE_LENGTH:
        return {'suggestions': [], 'indexed_values': 0, 'match_mode': 'none'}
    if column not in file_info.get('headers', []):
        raise ValueError('Selected column does not exist in this uploaded file.')
    source = Path(file_info['server_path']).resolve()
    upload_dir = (session_dir(session_id) / 'uploads').resolve()
    if source.parent != upload_dir or not source.is_file():
        raise ValueError('Uploaded file is missing or no longer belongs to this session.')
    cache = _cache_file(session_id, side, file_info, column)
    if not cache.is_file():
        with _INDEX_LOCK:
            if not cache.is_file():
                plan = engine.inspect_input(source, file_info.get('sheet') if source.suffix.casefold() != '.csv' else None)
                if column not in plan.headers:
                    raise ValueError('Column is no longer present in the selected worksheet.')
                _build_index(cache, plan, column)
    cap = min(10, max(1, int(limit)))
    with sqlite3.connect(f'file:{cache.as_posix()}?mode=ro', uri=True, timeout=10) as db:
        indexed_values = db.execute('SELECT count FROM metadata').fetchone()[0]
        if not search:
            found = [row[0] for row in db.execute(
                'SELECT value FROM suggestions ORDER BY folded LIMIT ?', (cap,))]
            mode = 'examples'
        else:
            found = [row[0] for row in db.execute(
                'SELECT value FROM suggestions WHERE folded >= ? AND folded < ? '
                'ORDER BY folded LIMIT ?',
                (search, search + chr(0x10ffff), cap))]
            mode = 'prefix'
            # This is intentionally a fallback, not the first search path:
            # large files retain the indexed prefix fast path.
            if len(found) < cap:
                existing = {item.casefold() for item in found}
                for (value,) in db.execute(
                    'SELECT value FROM suggestions WHERE instr(folded, ?) > 0 '
                    'ORDER BY folded LIMIT ?', (search, cap + len(existing))):
                    if value.casefold() not in existing:
                        found.append(value)
                        existing.add(value.casefold())
                    if len(found) >= cap:
                        break
                if any(not item.casefold().startswith(search) for item in found):
                    mode = 'contains'
    return {'suggestions': found, 'indexed_values': indexed_values, 'match_mode': mode}


def suggestions_for_file(session_id: str, side: str, file_info: dict,
                         column: str, prefix: str, limit: int = 10) -> list[str]:
    """Compatibility wrapper for other callers expecting only a list."""
    return lookup_suggestions_for_file(session_id, side, file_info, column, prefix, limit)['suggestions']
